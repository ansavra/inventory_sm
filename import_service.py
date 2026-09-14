"""
សេវានាំចូលទិន្នន័យ (Import Service)
អាន CSV / XLSX ទៅជា rows និងបង្កើតឯកសារគំរូ (Sample) សម្រាប់ទាញយក
"""
import csv
import io
import datetime
from typing import Any, List, Tuple, Dict


def _decode_text(raw: bytes) -> str:
    """Decode CSV bytes (គាំទ្រ UTF-8 with BOM, UTF-8, UTF-16 និង cp1252 ជាជម្រើសចុងក្រោយ)"""
    if raw.startswith(b'\xff\xfe') or raw.startswith(b'\xfe\xff'):
        return raw.decode('utf-16')
    for enc in ('utf-8-sig', 'utf-8'):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode('cp1252', errors='replace')


def parse_csv(raw: bytes) -> Tuple[List[Any], List[List[Any]]]:
    text = _decode_text(raw)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for r in reader]
    # រំលងជួរទទេ
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError("ឯកសារ CSV គ្មានទិន្នន័យ!")
    headers = rows[0]
    return headers, rows[1:]


def parse_xlsx(raw: bytes) -> Tuple[List[Any], List[List[Any]]]:
    try:
        import openpyxl
    except ImportError:
        raise ValueError("មិនអាចអាន Excel (.xlsx) បានទេ - សូមដំឡើង openpyxl (pip install openpyxl) ឬប្រើ CSV ជំនួស")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows: List[List[Any]] = []
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if any(v is not None and str(v).strip() for v in vals):
            rows.append(vals)
    wb.close()
    if not rows:
        raise ValueError("ឯកសារ Excel គ្មានទិន្នន័យ!")
    headers = rows[0]
    data = []
    for r in rows[1:]:
        # បម្លែង datetime ពី Excel ទៅជា string YYYY-MM-DD
        data.append([v.strftime('%Y-%m-%d') if isinstance(v, (datetime.date, datetime.datetime)) else v for v in r])
    return headers, data


def parse_upload(filename: str, raw: bytes) -> Tuple[List[Any], List[List[Any]]]:
    """ត្រឡប់ (headers, rows) ពីឯកសារ CSV ឬ XLSX"""
    name = (filename or '').lower()
    if name.endswith('.xlsx') or name.endswith('.xlsm') or raw[:2] == b'PK':
        return parse_xlsx(raw)
    if name.endswith('.xls'):
        raise ValueError("ទម្រង់ .xls ចាស់មិនគាំទ្រទេ - សូម Save As ជា .xlsx ឬ .csv")
    return parse_csv(raw)


def rows_to_csv(rows: List[List[Any]]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    for r in rows:
        writer.writerow(["" if v is None else v for v in r])
    return output.getvalue()


def rows_to_xlsx(rows: List[List[Any]]) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise RuntimeError("មិនអាចបង្កើត Excel បានទេ - សូមដំឡើង openpyxl ឬទាញយកជា CSV ជំនួស")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preview_rows(rows: List[Dict[str, Any]], limit: int = 20) -> List[Dict[str, Any]]:
    """ជួរខ្លះសម្រាប់បង្ហាញមុនពេលនាំចូល (Dry-run)"""
    out = []
    for r in rows[:limit]:
        out.append({k: ("" if v is None else str(v)) for k, v in r.items()})
    return out
